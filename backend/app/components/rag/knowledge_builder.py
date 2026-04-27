from __future__ import annotations

import hashlib
import logging
import tempfile
from pathlib import Path
from typing import AsyncIterator
from uuid import UUID

import httpx
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import (
    PyPDFLoader,
    TextLoader,
    WebBaseLoader,
)
from langchain_core.documents import Document
from qdrant_client import AsyncQdrantClient
from qdrant_client.models import (
    Distance,
    PointStruct,
    VectorParams,
    FilterSelector,
    Filter,
    FieldCondition,
    MatchValue,
)

from app.core.config import get_settings
from app.schemas.agent import RAGSpec

logger = logging.getLogger(__name__)
settings = get_settings()

EMBEDDING_DIM = 1536   # text-embedding-3-small


class KnowledgeBuilderService:
    """
    Gestiona el ciclo completo de construcción de una base de conocimiento RAG:
    1. Ingesta de fuentes (PDF, URLs, texto plano)
    2. Chunking configurable
    3. Generación de embeddings (OpenAI text-embedding-3-small)
    4. Almacenamiento en Qdrant (colección por agente)
    5. Retriever listo para usar en el AgentBuilder
    """

    def __init__(self):
        self._qdrant = AsyncQdrantClient(url=settings.qdrant_url)

    # ── Nombre de colección por agente ────────────────────────────────────────

    @staticmethod
    def collection_name(agent_id: str) -> str:
        return f"agent_{agent_id.replace('-', '_')}"

    # ── Ingesta principal ─────────────────────────────────────────────────────

    async def ingest(
        self,
        agent_id: str,
        rag_spec: RAGSpec,
        progress_cb: callable | None = None,
    ) -> dict:
        """
        Ingesta todas las fuentes del RAGSpec.
        Retorna estadísticas: { chunks_total, chunks_indexed, sources_ok, sources_failed }
        """
        collection = self.collection_name(agent_id)
        await self._ensure_collection(collection)

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=rag_spec.chunk_size,
            chunk_overlap=rag_spec.chunk_overlap,
            separators=["\n\n", "\n", ".", " ", ""],
        )

        stats = {"chunks_total": 0, "chunks_indexed": 0, "sources_ok": 0, "sources_failed": 0}

        for i, source in enumerate(rag_spec.sources):
            if progress_cb:
                await progress_cb(i, len(rag_spec.sources), source)
            try:
                docs = await self._load_source(source)
                chunks = splitter.split_documents(docs)
                stats["chunks_total"] += len(chunks)

                indexed = await self._embed_and_store(collection, chunks, source)
                stats["chunks_indexed"] += indexed
                stats["sources_ok"] += 1
                logger.info(f"[RAG] Ingested {indexed} chunks from {source}")
            except Exception as e:
                stats["sources_failed"] += 1
                logger.error(f"[RAG] Error ingesting {source}: {e}")

        return stats

    # ── Carga por tipo de fuente ──────────────────────────────────────────────

    async def _load_source(self, source: str) -> list[Document]:
        source = source.strip()

        if source.startswith("http://") or source.startswith("https://"):
            if source.lower().endswith(".pdf"):
                local_path = await self._download_file(source)
                return PyPDFLoader(str(local_path)).load()
            else:
                return WebBaseLoader([source]).load()

        path = Path(source)
        if path.exists():
            return self._load_file(path)

        # Tratar como texto directo
        return [Document(page_content=source, metadata={"source": "inline"})]

    def _load_file(self, path: Path) -> list[Document]:
        """Carga un archivo local según su extensión."""
        suffix = path.suffix.lower()

        if suffix == ".pdf":
            return PyPDFLoader(str(path)).load()

        if suffix in (".pptx", ".ppt"):
            from pptx import Presentation
            prs = Presentation(str(path))
            slides_text = []
            for i, slide in enumerate(prs.slides, 1):
                texts = [
                    shape.text.strip()
                    for shape in slide.shapes
                    if hasattr(shape, "text") and shape.text.strip()
                ]
                if texts:
                    slides_text.append(f"[Slide {i}]\n" + "\n".join(texts))
            content = "\n\n".join(slides_text)
            return [Document(page_content=content, metadata={"source": str(path)})]

        if suffix in (".docx", ".doc"):
            import docx
            doc = docx.Document(str(path))
            content = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
            return [Document(page_content=content, metadata={"source": str(path)})]

        if suffix in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            rows = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    line = "\t".join(str(c) for c in row if c is not None)
                    if line.strip():
                        rows.append(line)
            content = "\n".join(rows)
            return [Document(page_content=content, metadata={"source": str(path)})]

        if suffix == ".csv":
            content = path.read_text(encoding="utf-8", errors="replace")
            return [Document(page_content=content, metadata={"source": str(path)})]

        # Fallback: texto plano
        return TextLoader(str(path), encoding="utf-8", autodetect_encoding=True).load()

    async def _download_file(self, url: str) -> Path:
        tmp = Path(tempfile.gettempdir()) / hashlib.md5(url.encode()).hexdigest()
        async with httpx.AsyncClient(follow_redirects=True, timeout=30) as client:
            resp = await client.get(url)
            resp.raise_for_status()
            tmp.write_bytes(resp.content)
        return tmp

    # ── Embeddings y almacenamiento en Qdrant ─────────────────────────────────

    async def _embed_and_store(
        self,
        collection: str,
        chunks: list[Document],
        source: str,
    ) -> int:
        if not chunks:
            return 0

        texts = [c.page_content for c in chunks]
        vectors = await self._embed_texts(texts)

        points = [
            PointStruct(
                id=int(hashlib.md5(f"{source}_{i}".encode()).hexdigest()[:8], 16),
                vector=vector,
                payload={
                    "text":   texts[i],
                    "source": source,
                    "chunk":  i,
                    "metadata": chunks[i].metadata,
                },
            )
            for i, vector in enumerate(vectors)
        ]

        await self._qdrant.upsert(collection_name=collection, points=points)
        return len(points)

    async def _embed_texts(self, texts: list[str]) -> list[list[float]]:
        """Genera embeddings usando la misma conexión configurada en el builder.

        Reutiliza provider + api_key del builder config (Bóveda IA → Configuración del Builder).
        Para OpenRouter usa 'openai/text-embedding-3-small'; para OpenAI directo
        usa 'text-embedding-3-small'; Anthropic no soporta embeddings y devuelve error claro.
        """
        from openai import AsyncOpenAI
        from app.services.llm_client import _get_builder_config
        from app.runtime.llm import resolve_base_url

        provider, _, api_key = await _get_builder_config()

        if not api_key:
            raise ValueError(
                "No hay clave API configurada en el builder. "
                "Configurala en Bóveda IA → Configuración del Builder."
            )

        if provider == "anthropic":
            raise ValueError(
                "Anthropic no soporta embeddings nativos. "
                "Para usar RAG configurá el builder con OpenRouter u OpenAI."
            )

        # Nombre del modelo de embedding según provider
        embed_model = (
            "openai/text-embedding-3-small" if provider == "openrouter"
            else "text-embedding-3-small"
        )
        base_url = resolve_base_url(provider, None)

        client = AsyncOpenAI(api_key=api_key, base_url=base_url)

        # Batch de 100 textos por llamada
        all_vectors = []
        for i in range(0, len(texts), 100):
            batch = texts[i:i+100]
            resp = await client.embeddings.create(
                model=embed_model,
                input=batch,
            )
            all_vectors.extend([e.embedding for e in resp.data])
        return all_vectors

    # ── Retrieval ─────────────────────────────────────────────────────────────

    async def retrieve(
        self,
        agent_id: str,
        query: str,
        top_k: int = 4,
    ) -> list[dict]:
        """
        Busca los top_k fragmentos más relevantes para la query.
        Retorna lista de { text, source, score }.
        """
        collection = self.collection_name(agent_id)
        query_vector = (await self._embed_texts([query]))[0]

        results = await self._qdrant.search(
            collection_name=collection,
            query_vector=query_vector,
            limit=top_k,
            with_payload=True,
        )
        return [
            {
                "text":   r.payload["text"],
                "source": r.payload["source"],
                "score":  r.score,
            }
            for r in results
        ]

    async def retrieve_as_context(self, agent_id: str, query: str, top_k: int = 4) -> str:
        """Retorna los fragmentos formateados como contexto para el LLM."""
        chunks = await self.retrieve(agent_id, query, top_k)
        if not chunks:
            return ""
        lines = ["CONTEXTO RELEVANTE DE LA BASE DE CONOCIMIENTO:"]
        for i, c in enumerate(chunks, 1):
            lines.append(f"\n[{i}] (fuente: {c['source']})\n{c['text']}")
        return "\n".join(lines)

    # ── Gestión de colecciones ────────────────────────────────────────────────

    async def _ensure_collection(self, collection: str) -> None:
        existing = [c.name for c in (await self._qdrant.get_collections()).collections]
        if collection not in existing:
            await self._qdrant.create_collection(
                collection_name=collection,
                vectors_config=VectorParams(size=EMBEDDING_DIM, distance=Distance.COSINE),
            )
            logger.info(f"[RAG] Colección creada: {collection}")

    async def delete_collection(self, agent_id: str) -> None:
        collection = self.collection_name(agent_id)
        await self._qdrant.delete_collection(collection)

    async def get_sources(self, agent_id: str) -> list[str]:
        """Extrae todas las fuentes indexadas haciendo un scroll sin vectores."""
        collection = self.collection_name(agent_id)
        try:
            await self._ensure_collection(collection)
        except Exception:
            return []
        
        unique_sources = set()
        offset = None
        while True:
            records, next_offset = await self._qdrant.scroll(
                collection_name=collection,
                offset=offset,
                limit=100,
                with_payload=True,
                with_vectors=False
            )
            for r in records:
                if r.payload and "source" in r.payload:
                    unique_sources.add(r.payload["source"])
            offset = next_offset
            if offset is None:
                break
        return list(unique_sources)

    async def delete_source(self, agent_id: str, source: str) -> int:
        """Elimina todos los chunks de una fuente específica."""
        collection = self.collection_name(agent_id)
        result = await self._qdrant.delete(
            collection_name=collection,
            points_selector=FilterSelector(
                filter=Filter(
                    must=[FieldCondition(
                        key="source",
                        match=MatchValue(value=source)
                    )]
                )
            )
        )
        return getattr(result, "operation_id", 1)

    async def collection_stats(self, agent_id: str) -> dict:
        collection = self.collection_name(agent_id)
        try:
            info = await self._qdrant.get_collection(collection)
            return {
                "collection": collection,
                "vectors_count": info.vectors_count,
                "status": str(info.status),
            }
        except Exception:
            return {"collection": collection, "vectors_count": 0, "status": "not_found"}
